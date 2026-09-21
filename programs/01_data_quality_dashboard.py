#!/usr/bin/env python3
"""Audit global de qualité des données NetBox avec score et Excel coloré."""
from __future__ import annotations
import sys
from collections import Counter
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/"scripts"))
from common import *
log=setup("data_quality");args=parser("Qualité données NetBox").parse_args();rows=[]
for d in api().dcim.devices.filter(**filters(args)):
 r=device_row(d);checks={"serial":bool(r["serial"]),"primary_ip":bool(r["primary_ip"]),"site":bool(r["site"]),"role":bool(r["role"]),"model":bool(r["model"]),"status":bool(r["status"]),"location_or_rack":bool(r["location"] or r["rack"])};score=round(100*sum(checks.values())/len(checks));r.update({f"check_{k}":"PASS" if v else "FAIL" for k,v in checks.items()});r["score"]=score;r["severity"]="OK" if score==100 else "WARNING" if score>=70 else "CRITICAL";rows.append(r)
wb=Workbook();ws=wb.active;ws.title="Audit";headers=list(rows[0]) if rows else ["result"];ws.append(headers)
for r in rows:ws.append([r.get(h,"") for h in headers])
for c in ws[1]:c.font=Font(bold=True,color="FFFFFF");c.fill=PatternFill("solid",fgColor="1F4E78")
ws.freeze_panes="A2";summary=wb.create_sheet("Summary");summary.append(["Severity","Count"])
for k,v in Counter(r["severity"] for r in rows).items():summary.append([k,v])
path=OUTPUT/f"netbox_data_quality_{stamp()}.xlsx";wb.save(path);log.info("%s",path)

