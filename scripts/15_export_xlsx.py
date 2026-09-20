#!/usr/bin/env python3
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill
from openpyxl.worksheet.table import Table,TableStyleInfo
from common import *
log=setup("export_xlsx");nb=api();datasets={"Devices":[device_row(d) for d in nb.dcim.devices.all()],"Interfaces":[]}
for d in nb.dcim.devices.all():
 for i in nb.dcim.interfaces.filter(device_id=d.id):datasets["Interfaces"].append({"site":val(d,"site"),"rack":val(d,"rack"),"location":val(d,"location"),"device":d.name,"interface":i.name,"description":i.description,"enabled":i.enabled,"connected":bool(i.connected_endpoints)})
wb=Workbook();wb.remove(wb.active)
for title,rows in datasets.items():
 ws=wb.create_sheet(title);headers=sorted({k for r in rows for k in r}) or ["result"];ws.append(headers)
 for r in rows:ws.append([r.get(h,"") for h in headers])
 for c in ws[1]:c.font=Font(bold=True,color="FFFFFF");c.fill=PatternFill("solid",fgColor="1F4E78")
 ws.freeze_panes="A2";ws.auto_filter.ref=ws.dimensions
 for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=min(40,max(12,max(len(str(c.value or "")) for c in col)+2))
 if rows:
  t=Table(displayName=f"Table{title}",ref=ws.dimensions);t.tableStyleInfo=TableStyleInfo(name="TableStyleMedium2",showRowStripes=True);ws.add_table(t)
path=OUTPUT/f"netbox_export_{stamp()}.xlsx";wb.save(path);log.info("%s",path)
